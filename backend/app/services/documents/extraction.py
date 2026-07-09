import io

import openpyxl
import pdfplumber
from docx import Document as DocxDocument
from pypdf import PdfReader

_DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def extract_text(data: bytes, mime_type: str, filename: str) -> str:
    if mime_type == "application/pdf":
        return _extract_pdf(data)
    if mime_type == _DOCX_MIME:
        return _extract_docx(data)
    if mime_type == _XLSX_MIME:
        return _extract_xlsx(data)
    if mime_type.startswith("text/"):
        return data.decode("utf-8", errors="replace")
    raise ValueError(f"Unsupported mime type for extraction: {mime_type} ({filename})")


def _extract_pdf(data: bytes) -> str:
    text_parts: list[str] = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            text_parts.append(page.extract_text() or "")
    text = "\n\n".join(text_parts).strip()
    if text:
        return text

    # Fallback for PDFs pdfplumber can't get text out of (e.g. unusual encodings).
    reader = PdfReader(io.BytesIO(data))
    return "\n\n".join((page.extract_text() or "") for page in reader.pages).strip()


def _extract_docx(data: bytes) -> str:
    doc = DocxDocument(io.BytesIO(data))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def _extract_xlsx(data: bytes) -> str:
    workbook = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    lines: list[str] = []
    for sheet in workbook.worksheets:
        lines.append(f"# Sheet: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(cell) for cell in row if cell is not None]
            if cells:
                lines.append("\t".join(cells))
    return "\n".join(lines)
